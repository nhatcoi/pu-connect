#!/usr/bin/env python3
"""
Phenikaa PowerPoint Presentation Generator
Uses the official Phenikaa University slide template (navy, orange, hex pattern, logos).
"""

import sys
import json
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x22, 0x37, 0x71)
ORANGE = RGBColor(0xF2, 0x65, 0x22)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(0xEE, 0xF1, 0xF8)
FONT = "Times New Roman"
X0, X1, Y0, Y1 = 560000, 11630000, 1050000, 6150000

class PhenikaaDeck:
    def __init__(self, template_path):
        self.p = Presentation(template_path)
        ids = self.p.slides._sldIdLst
        for s in list(ids):
            self.p.part.drop_rel(s.rId)
            ids.remove(s)
        self.L = {l.name: l for l in self.p.slide_layouts}

    def _slide(self, layout, title=None):
        target_layout = self.L.get(layout) or self.L.get('Title and Content') or self.p.slide_layouts[0]
        s = self.p.slides.add_slide(target_layout)
        if title is not None and s.shapes.title is not None:
            s.shapes.title.text = title
            if s.shapes.title.text_frame and s.shapes.title.text_frame.paragraphs:
                p = s.shapes.title.text_frame.paragraphs[0]
                if p.runs:
                    p.runs[0].font.name = FONT
                    p.runs[0].font.color.rgb = NAVY
                    p.runs[0].font.bold = True
        return s

    def _text(self, s, x, y, w, h, paras, size=20, color=INK, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullet=False, italic=False):
        tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        paras = [paras] if isinstance(paras, str) else paras
        for i, t in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(8)
            r = para.add_run()
            r.text = ("•  " + str(t)) if bullet else str(t)
            f = r.font
            f.name = FONT
            f.size = Pt(size)
            f.bold = bold
            f.italic = italic
            f.color.rgb = color
        return tb

    def cover(self, title, subtitle=""):
        s = self._slide("Title Slide")
        if s.shapes.title:
            t = s.shapes.title
            t.text = title.upper()
            t.left, t.top, t.width, t.height = Emu(493486), Emu(3782987), Emu(11219542), Emu(906811)
            if t.text_frame and t.text_frame.paragraphs and t.text_frame.paragraphs[0].runs:
                r = t.text_frame.paragraphs[0].runs[0]
                r.font.name = FONT
                r.font.size = Pt(36)
                r.font.bold = True
                r.font.color.rgb = WHITE
        if len(s.placeholders) > 1 and subtitle:
            sub = s.placeholders[1]
            sub.text = subtitle
            sub.left, sub.top, sub.width, sub.height = Emu(493486), Emu(4750000), Emu(11219542), Emu(900000)
            for para in sub.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for r in para.runs:
                    r.font.name = FONT
                    r.font.size = Pt(20)
                    r.font.color.rgb = WHITE
        return s

    def closing(self, text="Xin cảm ơn!", sub=""):
        s = self._slide("Title Slide")
        if s.shapes.title:
            t = s.shapes.title
            t.text = text
            t.left, t.top, t.width, t.height = Emu(493486), Emu(3782987), Emu(11219542), Emu(906811)
            if t.text_frame and t.text_frame.paragraphs and t.text_frame.paragraphs[0].runs:
                r = t.text_frame.paragraphs[0].runs[0]
                r.font.name = FONT
                r.font.size = Pt(40)
                r.font.bold = True
                r.font.color.rgb = WHITE
        if len(s.placeholders) > 1 and sub:
            sub_ph = s.placeholders[1]
            sub_ph.text = sub
            sub_ph.left, sub_ph.top, sub_ph.width, sub_ph.height = Emu(493486), Emu(4750000), Emu(11219542), Emu(900000)
            for para in sub_ph.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for r in para.runs:
                    r.font.name = FONT
                    r.font.size = Pt(20)
                    r.font.color.rgb = WHITE
        return s

    def section(self, title, kicker="CHỦ ĐỀ"):
        s = self._slide("Section Header", title)
        if s.shapes.title and s.shapes.title.text_frame.paragraphs:
            for r in s.shapes.title.text_frame.paragraphs[0].runs:
                r.font.name = FONT
                r.font.size = Pt(44)
                r.font.bold = True
                r.font.color.rgb = NAVY
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 1:
                ph.text = kicker
                if ph.text_frame.paragraphs and ph.text_frame.paragraphs[0].runs:
                    r = ph.text_frame.paragraphs[0].runs[0]
                    r.font.name = FONT
                    r.font.size = Pt(18)
                    r.font.bold = True
                    r.font.color.rgb = ORANGE
        return s

    def bullets(self, title, items, size=22):
        s = self._slide("Title and Content", title)
        self._text(s, X0, Y0, 9500000, Y1 - Y0, items, size=size, bullet=True)
        return s

    def two_col(self, title, left_head, left_items, right_head, right_items, size=19):
        s = self._slide("Title and Content", title)
        w = (X1 - X0 - 400000) // 2
        cols = [(left_head, left_items), (right_head, right_items)]
        for i, (head, items) in enumerate(cols):
            x = X0 + i * (w + 400000)
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(Y0), Emu(w), Emu(520000))
            box.fill.solid()
            box.fill.fore_color.rgb = NAVY
            box.line.fill.background()
            tf = box.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = head or ""
            r.font.name = FONT
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = WHITE
            self._text(s, x, Y0 + 620000, w, Y1 - Y0 - 620000, items or [], size=size, bullet=True)
        return s

    def stats(self, title, stats):
        s = self._slide("Title and Content", title)
        n = min(4, max(1, len(stats)))
        gap = 300000
        w = (X1 - X0 - gap * (n - 1)) // n
        for i, st in enumerate(stats[:n]):
            big = str(st.get("value", ""))
            label = str(st.get("label", ""))
            x = X0 + i * (w + gap)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(2000000), Emu(w), Emu(2800000))
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT_BG
            card.line.fill.background()
            card.shadow.inherit = False
            self._text(s, x, 2200000, w, 1200000, big, size=44, color=ORANGE, bold=True,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, x + 100000, 3500000, w - 200000, 1100000, label, size=18, color=NAVY,
                       align=PP_ALIGN.CENTER, bold=True)
        return s

    def table(self, title, header, rows, size=16):
        s = self._slide("Title and Content", title)
        num_rows = len(rows) + 1
        num_cols = len(header)
        h = min(Y1 - Y0, 480000 * num_rows)
        gt = s.shapes.add_table(num_rows, num_cols, Emu(X0), Emu(Y0), Emu(X1 - X0), Emu(h)).table
        for c, t in enumerate(header):
            cell = gt.cell(0, c)
            cell.text = str(t)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            for r in cell.text_frame.paragraphs[0].runs:
                r.font.name = FONT
                r.font.size = Pt(size)
                r.font.bold = True
                r.font.color.rgb = WHITE
        for i, row in enumerate(rows, 1):
            for c, t in enumerate(row[:num_cols]):
                cell = gt.cell(i, c)
                cell.text = str(t)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT_BG
                for r in cell.text_frame.paragraphs[0].runs:
                    r.font.name = FONT
                    r.font.size = Pt(size)
                    r.font.color.rgb = INK
        return s

    def notes(self, slide, text):
        if text and hasattr(slide, 'notes_slide'):
            slide.notes_slide.notes_text_frame.text = text

    def save(self, output_path):
        self.p.save(output_path)


def generate_deck(data, template_path, output_path):
    deck = PhenikaaDeck(template_path)

    # 1. Slide bìa
    title = data.get("title", "BÀI THUYẾT TRÌNH")
    subtitle = data.get("subtitle", "Trường Đại học Phenikaa")
    deck.cover(title, subtitle)

    # 2. Các slide nội dung
    slides = data.get("slides", [])
    has_closing = False

    for item in slides:
        layout = item.get("layout", "bullets")
        s_title = item.get("title", "")
        notes = item.get("notes")
        slide_obj = None

        if layout == "section":
            kicker = item.get("kicker", "PHẦN")
            slide_obj = deck.section(s_title, kicker)
        elif layout == "two_col":
            left_head = item.get("leftHead", "Phần 1")
            left_items = item.get("leftBullets", [])
            right_head = item.get("rightHead", "Phần 2")
            right_items = item.get("rightBullets", [])
            slide_obj = deck.two_col(s_title, left_head, left_items, right_head, right_items)
        elif layout == "stats":
            stats = item.get("stats", [])
            slide_obj = deck.stats(s_title, stats)
        elif layout == "table":
            tbl = item.get("table", {})
            header = tbl.get("headers", [])
            rows = tbl.get("rows", [])
            slide_obj = deck.table(s_title, header, rows)
        elif layout == "closing":
            has_closing = True
            c_text = item.get("closingText") or s_title or "Xin cảm ơn!"
            c_sub = item.get("closingSub", "Q&A · Xin mời quý Thầy Cô và các bạn đặt câu hỏi")
            slide_obj = deck.closing(c_text, c_sub)
        else: # default bullets
            bullets = item.get("bullets", [])
            slide_obj = deck.bullets(s_title, bullets)

        if notes and slide_obj:
            deck.notes(slide_obj, notes)

    # Tự động thêm slide closing nếu chưa có
    if not has_closing:
        deck.closing("Xin cảm ơn!", "Q&A · Xin mời quý Thầy Cô và các bạn đặt câu hỏi")

    deck.save(output_path)


def main():
    if len(sys.argv) < 3:
        print("Usage: generate_phenikaa_deck.py <template.pptx> <output.pptx> [input.json]", file=sys.stderr)
        sys.exit(1)

    template_path = sys.argv[1]
    output_path = sys.argv[2]

    if len(sys.argv) >= 4 and sys.argv[3] != "-":
        with open(sys.argv[3], "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = json.load(sys.stdin)

    generate_deck(data, template_path, output_path)
    print(f"OK:{output_path}")

if __name__ == "__main__":
    main()
